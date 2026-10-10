"""Gera apresentacao/EdgeHealth_Apresentacao_Final.pptx (16 slides 16:9, com notas do apresentador).

Uso (numa venv fora do repositório, só com python-pptx, que já traz o Pillow):
    python -m venv %TEMP%\\pptvenv
    %TEMP%\\pptvenv\\Scripts\\python -m pip install python-pptx==1.0.2
    %TEMP%\\pptvenv\\Scripts\\python apresentacao\\gerar_apresentacao.py

As imagens vêm do repositório (docs/img/telas, docs/img/fluxo-*.png, docs/img/diagrama-classes.png).
O logo (frontend/public/logo.svg) é convertido para PNG pelo Chrome, com o playwright-core do frontend
(rode `npm install` em frontend/ antes). Os números citados (testes, funcionalidades) são os da
auditoria de outubro de 2026: confira-os no README antes de regerar.
"""
import subprocess
import tempfile
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'apresentacao' / 'EdgeHealth_Apresentacao_Final.pptx'
IMG = ROOT / 'docs' / 'img'
SITE = 'https://marcelodomingos.pythonanywhere.com'
REPO = 'github.com/MarceloRodriguesDoNascimento/EdgeHealth'

TEAL = RGBColor(0x08, 0x7F, 0x8C)
TEAL_DARK = RGBColor(0x06, 0x5F, 0x69)
MINT = RGBColor(0xBD, 0xF2, 0xEA)
NAVY = RGBColor(0x12, 0x23, 0x3A)
INK = RGBColor(0x1F, 0x2A, 0x37)
MUTED = RGBColor(0x5B, 0x6B, 0x7B)
LIGHT = RGBColor(0xF1, 0xF6, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xC6, 0x28, 0x28)
FONT = 'Segoe UI'
MONO = 'Consolas'

W, H = Inches(13.333), Inches(7.5)

MD, MR, ER, FE, MA, JL = ('Marcelo Domingos', 'Marcelo Rodrigues Alves do Nascimento', 'Erick Daniel Coelho E Silva',
                          'Felipe Barbosa Poeiras', 'Matheus Brito Vaz Bernardes', 'João Lucas Santos Batista')
TEAM = [(MD, '22400362'), (MR, '12400815'), (ER, '12401188'), (FE, '12402320'), (MA, '12502391'), (JL, '12401617')]


# ---------------------------------------------------------------------------------------------
# Images
# ---------------------------------------------------------------------------------------------

LOGO_JS = r"""
import {createRequire} from 'node:module';
import {readFileSync} from 'node:fs';
const require = createRequire(process.cwd() + '/package.json');
const {chromium} = require('playwright-core');
const [svg, out, exe] = process.argv.slice(1);
const browser = await chromium.launch({executablePath: exe, headless: true});
const page = await browser.newPage({viewport: {width: 512, height: 512}});
await page.setContent(`<body style="margin:0">${readFileSync(svg, 'utf8').replace('<svg ', '<svg width="512" height="512" ')}</body>`);
await page.screenshot({path: out, omitBackground: true, clip: {x: 0, y: 0, width: 512, height: 512}});
await browser.close();
"""


def find_browser():
    for path in ('C:/Program Files/Google/Chrome/Application/chrome.exe',
                 'C:/Program Files (x86)/Microsoft/Edge/Application/msedge.exe',
                 'C:/Program Files/Microsoft/Edge/Application/msedge.exe',
                 '/usr/bin/google-chrome', '/usr/bin/chromium'):
        if Path(path).exists():
            return path
    raise SystemExit('Chrome ou Edge não encontrado para converter o logo.')


def prepare_images(tmp):
    logo = tmp / 'logo.png'
    subprocess.run(['node', '--input-type=module', '-e', LOGO_JS, str(ROOT / 'frontend/public/logo.svg'), str(logo), find_browser()],
                   cwd=ROOT / 'frontend', check=True)
    crops = {  # name -> (source, height of the top crop in px; None = whole image)
        'dashboard': ('telas/visao-da-rede.png', 860), 'dispositivo': ('telas/modal-dispositivo.png', None),
        'coletores': ('telas/coletores.png', None), 'falha': ('telas/falha.png', 1000),
        'historico': ('telas/historico.png', None), 'metricas': ('telas/modal-metricas.png', None)}
    images = {'logo': logo}
    for name, (source, height) in crops.items():
        im = Image.open(IMG / source).convert('RGB')
        if height:
            im = im.crop((0, 0, im.width, min(height, im.height)))
        images[name] = tmp / f'{name}.png'
        im.save(images[name])
    return images


# ---------------------------------------------------------------------------------------------
# Drawing helpers
# ---------------------------------------------------------------------------------------------

def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def no_line(shape):
    shape.line.fill.background()


def text(slide, x, y, w, h, content, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT, font=FONT,
         anchor=MSO_ANCHOR.TOP, spacing=1.1):
    """Text box; content is a string or a list of paragraphs (str or (str, size, color, bold))."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    paragraphs = content if isinstance(content, list) else [content]
    for i, item in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        value, psize, pcolor, pbold = (item if isinstance(item, tuple) else (item, size, color, bold))
        p.alignment = align
        p.line_spacing = spacing
        run = p.add_run()
        run.text = value
        run.font.size, run.font.color.rgb, run.font.bold, run.font.name = Pt(psize), pcolor, pbold, font
    return box


def bullets(slide, x, y, w, h, items, size=22, color=INK, gap=10):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        pPr = p._p.get_or_add_pPr()
        pPr.set('marL', str(Inches(0.35)))
        pPr.set('indent', str(-Inches(0.35)))
        bu = pPr.makeelement(qn('a:buChar'), {'char': '•'})
        pPr.append(bu)
        run = p.add_run()
        run.text = item
        run.font.size, run.font.color.rgb, run.font.name = Pt(size), color, FONT
    return box


def box(slide, x, y, w, h, title=None, body=None, color=LIGHT, title_color=TEAL_DARK, body_size=20, title_size=22,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, border=None, body_color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        margin=0.18):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    fill(s, color)
    if border:
        s.line.color.rgb = border
        s.line.width = Pt(1.5)
    else:
        no_line(s)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right'):
        setattr(tf, m, Inches(margin))
    tf.margin_top = tf.margin_bottom = Inches(0.12)
    first = True
    for value, size, col, bold in ([(title, title_size, title_color, True)] if title else []) + \
            [(line, body_size, body_color, False) for line in ([body] if isinstance(body, str) else body or [])]:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = value
        run.font.size, run.font.color.rgb, run.font.bold, run.font.name = Pt(size), col, bold, FONT
    return s


def arrow(slide, x1, y1, x2, y2, color=TEAL, width=2.5, label=None, label_size=16):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))
    if label:
        text(slide, min(x1, x2) - Inches(0.6), (y1 + y2) // 2 - Inches(0.42), abs(x2 - x1) + Inches(1.2), Inches(0.4),
             label, size=label_size, color=MUTED, align=PP_ALIGN.CENTER)
    return c


def picture(slide, path, x, y, max_w, max_h, border=True):
    """Image scaled to fit inside the box, centred horizontally, top-aligned."""
    with Image.open(path) as im:
        ratio = im.width / im.height
    w, h = max_w, int(max_w / ratio)
    if h > max_h:
        h, w = max_h, int(max_h * ratio)
    pic = slide.shapes.add_picture(str(path), x + (max_w - w) // 2, y, w, h)
    if border:
        pic.line.color.rgb = RGBColor(0xD5, 0xDE, 0xE3)
        pic.line.width = Pt(1)
    return pic


def base(prs, title, kicker=None, number=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = WHITE
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.42), Inches(0.12), Inches(0.62))
    fill(bar, TEAL)
    no_line(bar)
    if kicker:
        text(slide, Inches(0.75), Inches(0.22), Inches(11), Inches(0.35), kicker.upper(), size=14, color=TEAL, bold=True)
    text(slide, Inches(0.75), Inches(0.45), Inches(12), Inches(0.75), title, size=34, color=NAVY, bold=True)
    footer = f'EdgeHealth · Projeto de Software · PS-3B1/2026' + (f'   {number}' if number else '')
    text(slide, Inches(0.5), Inches(7.05), Inches(12.33), Inches(0.3), footer, size=11, color=MUTED, align=PP_ALIGN.RIGHT)
    return slide


def notes(slide, speaker, seconds, lines):
    minutes, rest = divmod(seconds, 60)
    tf = slide.notes_slide.notes_text_frame
    tf.text = f'Quem fala: {speaker} · Tempo sugerido: {minutes}:{rest:02d}'
    for line in lines:
        tf.add_paragraph().text = line


# ---------------------------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------------------------

def cover(prs, img):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = TEAL
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.55), 0, Inches(5.783), H)
    fill(panel, NAVY)
    no_line(panel)
    slide.shapes.add_picture(str(img['logo']), Inches(0.8), Inches(0.9), Inches(1.5), Inches(1.5))
    text(slide, Inches(0.8), Inches(2.6), Inches(6.5), Inches(1.2), 'EdgeHealth', size=60, color=WHITE, bold=True)
    text(slide, Inches(0.8), Inches(3.75), Inches(6.5), Inches(2.0),
         'Monitoramento e diagnóstico de redes para pequenas e médias empresas: detecta a falha, explica a causa provável e estima o prejuízo.',
         size=24, color=MINT)
    text(slide, Inches(0.8), Inches(5.9), Inches(6.5), Inches(0.9),
         ['[PS-3B1/2026] Projeto de Software', ('Colégio Cotemig · apresentação final · 07/11/2026', 18, MINT, False)],
         size=20, color=WHITE, bold=True)
    text(slide, Inches(7.9), Inches(0.75), Inches(5.2), Inches(0.5), 'EQUIPE', size=16, color=MINT, bold=True)
    y = Inches(1.3)
    for name, reg in TEAM:
        text(slide, Inches(7.9), y, Inches(5.3), Inches(0.9), [(name, 19, WHITE, True), (f'Matrícula {reg}', 15, MINT, False)], spacing=1.0)
        y += Inches(0.93)
    notes(slide, MD, 25, [
        'Bom dia! Nós somos a equipe do EdgeHealth.',
        'Eu sou o Marcelo Domingos e comigo estão o Marcelo Rodrigues, o Erick, o Felipe, o Matheus e o João Lucas.',
        'Em poucas palavras: o EdgeHealth monitora a rede de pequenas empresas, mostra a causa provável de uma falha e quanto ela custou.'])


def problem(prs):
    slide = base(prs, 'O problema e para quem é', 'Problema e público-alvo', 2)
    items = [('A rede cai e a empresa só descobre quando alguém reclama.', '1'),
             ('Ninguém sabe dizer a causa provável da falha.', '2'),
             ('Ninguém sabe quanto a parada custou em dinheiro.', '3')]
    y = Inches(1.55)
    for sentence, n in items:
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.75), y, Inches(0.75), Inches(0.75))
        fill(circle, TEAL)
        no_line(circle)
        circle.text_frame.text = n
        p = circle.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.runs[0].font.size, p.runs[0].font.bold, p.runs[0].font.color.rgb, p.runs[0].font.name = Pt(24), True, WHITE, FONT
        text(slide, Inches(1.75), y, Inches(5.6), Inches(1.4), sentence, size=24, color=INK, anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(1.7)
    box(slide, Inches(7.75), Inches(1.55), Inches(5.08), Inches(5.0), 'Público-alvo', [
        'Pequenas e médias empresas sem equipe de TI dedicada.',
        'O técnico ou prestador de TI que atende essas empresas.',
        'O gestor, que precisa entender a falha sem termos técnicos.'], color=LIGHT, body_size=22, title_size=26)
    notes(slide, MD, 45, [
        'Numa empresa pequena, quando a internet ou a impressora param, ninguém fica sabendo até alguém reclamar.',
        'E mesmo quando descobrem, ninguém sabe dizer o motivo provável, nem quanto aquela parada custou.',
        'Nosso público são essas pequenas e médias empresas sem equipe de TI, o técnico que atende elas e o gestor, que quer uma explicação simples.'])


def solution(prs):
    slide = base(prs, 'Solução: medir, diagnosticar, estimar e explicar', 'Como funciona', 3)
    text(slide, Inches(0.75), Inches(1.35), Inches(12), Inches(0.9),
         'O EdgeHealth monitora a rede da empresa, aponta a causa provável de cada falha e mostra quanto ela custou.',
         size=24, color=TEAL_DARK, bold=True)
    steps = [('Coletor', 'mede os equipamentos da rede (ICMP)'), ('Servidor', 'detecta a falha e calcula a severidade'),
             ('Diagnóstico', 'causa provável e recomendações'), ('Prejuízo', 'horas paradas × pessoas × custo'),
             ('IA', 'explica em português simples')]
    x, w, gap = Inches(0.6), Inches(2.25), Inches(0.27)
    for i, (title, body) in enumerate(steps):
        box(slide, x, Inches(2.75), w, Inches(2.6), title, body, color=TEAL if i in (0, 4) else LIGHT,
            title_color=WHITE if i in (0, 4) else TEAL_DARK, body_color=WHITE if i in (0, 4) else INK,
            body_size=20, title_size=24, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.06)
        if i < len(steps) - 1:
            arrow(slide, x + w + Inches(0.02), Inches(4.05), x + w + gap - Inches(0.02), Inches(4.05))
        x += w + gap
    text(slide, Inches(0.75), Inches(5.75), Inches(12), Inches(1.0),
         'Tudo calculado a partir de medições reais: não há dados simulados no produto.', size=22, color=MUTED)
    notes(slide, MD, 50, [
        'Nossa solução funciona em cinco passos.',
        'Um coletor instalado na rede da empresa mede os equipamentos e manda as medições pro servidor.',
        'O servidor detecta a falha, calcula a severidade e faz um diagnóstico por regras, com recomendações.',
        'Ele também estima o prejuízo, e a IA explica tudo em linguagem simples pro gestor.',
        'Agora o Marcelo Rodrigues vai mostrar os principais casos de uso.'])


def split_name(name):
    """'Classe.metodo' in two lines, so long identifiers stay legible at a large size."""
    cls, dot, method = name.partition('.')
    return [cls, dot + method] if dot and len(name) > 26 else [name]


def use_case(prs, number, slide_no, title, kind, what, image, flow, chain, speaker, seconds, lines):
    slide = base(prs, title, f'Caso de uso {number} · {kind}', slide_no)
    text(slide, Inches(0.75), Inches(1.22), Inches(12), Inches(0.5), what, size=20, color=MUTED)
    picture(slide, image, Inches(0.5), Inches(1.9), Inches(7.55), Inches(4.95))
    x, y, w = Inches(8.3), Inches(1.9), Inches(3.3)
    step = Inches(1.25)
    for i, (layer, name) in enumerate(chain):
        b = box(slide, x, y, w, Inches(1.05), layer.upper(), split_name(name), color=LIGHT if layer != 'Tela' else MINT,
                title_size=12, body_size=15, title_color=TEAL)
        b.text_frame.margin_top = b.text_frame.margin_bottom = Inches(0.04)
        if i < len(chain) - 1:
            arrow(slide, x + w // 2, y + Inches(1.05), x + w // 2, y + step, width=2)
        y += step
    pic = picture(slide, IMG / flow, Inches(11.8), Inches(1.9), Inches(1.1), Inches(4.4))
    text(slide, Inches(11.55), pic.top + pic.height + Inches(0.05), Inches(1.6), Inches(0.6), f'Fluxograma {number}',
         size=12, color=MUTED, align=PP_ALIGN.CENTER)
    notes(slide, speaker, seconds, lines)


def architecture(prs):
    slide = base(prs, 'Arquitetura cliente-servidor', 'Arquitetura', 9)
    box(slide, Inches(0.6), Inches(1.6), Inches(3.1), Inches(1.9), 'Cliente (frontend)',
        ['Navegador', 'Vite + JavaScript', 'Chart.js'], color=MINT, body_size=20)
    api = box(slide, Inches(5.0), Inches(1.35), Inches(4.1), Inches(4.6), 'Servidor: API Flask', None, color=LIGHT, border=TEAL)
    for i, (layer, color) in enumerate([('Controllers', WHITE), ('Services', WHITE), ('Models + Repositories', WHITE)]):
        box(slide, Inches(5.3), Inches(2.05) + i * Inches(1.25), Inches(3.5), Inches(0.85), None, layer, color=color,
            border=TEAL, body_size=20, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < 2:
            arrow(slide, Inches(7.05), Inches(2.9) + i * Inches(1.25), Inches(7.05), Inches(3.3) + i * Inches(1.25), width=2)
    box(slide, Inches(10.55), Inches(3.95), Inches(2.3), Inches(1.6), None, ['SQLite', '13 tabelas'], color=NAVY,
        body_color=WHITE, shape=MSO_SHAPE.CAN, body_size=20, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    box(slide, Inches(10.4), Inches(1.6), Inches(2.6), Inches(1.3), 'Google Gemini', 'IA como serviço', color=LIGHT,
        body_size=18, title_size=20, align=PP_ALIGN.CENTER)
    box(slide, Inches(0.6), Inches(4.3), Inches(3.1), Inches(1.65), 'Coletor', ['Na rede da empresa', 'mede por ICMP'],
        color=TEAL, title_color=WHITE, body_color=WHITE, body_size=20)
    arrow(slide, Inches(3.75), Inches(2.55), Inches(4.95), Inches(2.55), label='HTTP / JSON')
    arrow(slide, Inches(3.75), Inches(5.1), Inches(4.95), Inches(5.1), label='HTTPS')
    arrow(slide, Inches(9.15), Inches(4.75), Inches(10.5), Inches(4.75))
    arrow(slide, Inches(9.15), Inches(2.25), Inches(10.35), Inches(2.25), label='só dados técnicos', label_size=14)
    text(slide, Inches(0.6), Inches(6.2), Inches(12.2), Inches(0.8),
         'Em desenvolvimento: Vite em :5173 e Flask em :5000, servidores separados. Em produção, o Flask serve o build do frontend e a API (PythonAnywhere).',
         size=20, color=MUTED)
    notes(slide, FE, 40, [
        'O sistema é cliente-servidor. O frontend e a API são aplicações separadas.',
        'Em desenvolvimento rodam em dois servidores: o Vite na porta 5173 e o Flask na 5000. O frontend só conversa com a API por HTTP, com JSON.',
        'O coletor é outro cliente da mesma API, e a API é a única que fala com o banco e com o Gemini.',
        'Em produção o Flask serve o build do frontend junto com a API, no PythonAnywhere.',
        'E pra construir isso, usamos estas tecnologias.'])


def directories(prs):
    slide = base(prs, 'API em camadas: os diretórios', 'Organização do código', 11)
    rows = [('backend/', '', NAVY),
            ('├── controllers/', '12 classes, uma por recurso + rotas.py (39 rotas)', TEAL_DARK),
            ('├── services/', '68 classes, um caso de uso cada (executar)', TEAL_DARK),
            ('├── models/', '13 entidades; BaseModel com o CRUD', TEAL_DARK),
            ('├── repositories/', '13 classes, só consultas SQL especiais', TEAL_DARK),
            ('├── database/', 'create_database.sql', TEAL_DARK),
            ('└── app/', 'configuração, segurança e validação', TEAL_DARK),
            ('frontend/src/pages/', '9 telas; services/api.js faz o fetch', NAVY),
            ('collector/', 'coletor remoto (Python e .exe)', NAVY)]
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.4), Inches(12.1), Inches(4.6))
    fill(panel, LIGHT)
    no_line(panel)
    panel.adjustments[0] = 0.04
    y = Inches(1.6)
    for path, desc, color in rows:
        text(slide, Inches(0.9), y, Inches(4.3), Inches(0.5), path, size=22, color=color, bold=True, font=MONO)
        text(slide, Inches(5.2), y + Inches(0.02), Inches(7.4), Inches(0.5), desc, size=20, color=INK)
        y += Inches(0.48)
    text(slide, Inches(0.6), Inches(6.2), Inches(12.2), Inches(0.8),
         'Tela → Controller → Service → Model / Repository → Banco. Um teste automático falha se uma camada pular a outra.',
         size=20, color=MUTED)
    notes(slide, MA, 45, [
        'Dentro do backend, a API é organizada em camadas, cada uma numa pasta.',
        'Os controllers só recebem a requisição e chamam um service. São 12, um por recurso.',
        'Os services têm a regra de negócio: cada classe é um caso de uso, com o método executar. São 68.',
        'Os models fazem o CRUD, herdando de uma classe base, e os repositories têm só as consultas SQL especiais.',
        'O test_architecture falha se um controller acessar o banco ou chamar mais de um service.'])


def class_diagram(prs):
    slide = base(prs, 'Diagrama de classes do domínio', 'Modelagem', 12)
    picture(slide, IMG / 'diagrama-classes.png', Inches(0.5), Inches(1.3), Inches(6.2), Inches(5.6))
    bullets(slide, Inches(7.0), Inches(1.6), Inches(5.8), Inches(5.0), [
        '13 entidades, com os atributos principais.',
        'Relações com cardinalidade: uma Empresa tem vários Dispositivos; um Dispositivo tem várias Falhas.',
        'Cada Falha tem um Impacto e um Diagnóstico com Recomendações.',
        'Gerado a partir dos Models: um teste falha se o diagrama ficar diferente do código.'], size=22)
    notes(slide, MA, 30, [
        'Esse é o diagrama de classes do domínio. A empresa é o centro: ela tem usuários, coletores e dispositivos.',
        'Cada dispositivo gera métricas e pode ter falhas; cada falha tem impacto e diagnóstico, e o diagnóstico aponta recomendações.',
        'Ele é gerado a partir dos models e tem teste, então sempre corresponde ao sistema.'])


def database(prs):
    slide = base(prs, 'Banco de dados e SQL nos Repositories', 'Persistência', 13)
    bullets(slide, Inches(0.6), Inches(1.5), Inches(5.3), Inches(5.0), [
        'SQLite, com o esquema versionado em migrations (Alembic).',
        'database/create_database.sql cria o banco inteiro.',
        'CRUD pelo ORM (SQLAlchemy) nos Models.',
        'O SQLite não tem stored procedures: as consultas especiais ficam nos Repositories.'], size=22)
    rows = [('Consulta', 'SQL'),
            ('FalhaRepository.ranking_dispositivos', 'JOIN + GROUP BY + ORDER BY + LIMIT'),
            ('RankingCustoRepository.top_dispositivos', 'CTE + LEFT JOIN + ORDER BY'),
            ('DashboardRepository.contagem_por_status', 'COALESCE + GROUP BY'),
            ('FalhaRepository.historico', 'JOIN + WHERE + LIMIT/OFFSET')]
    table = slide.shapes.add_table(len(rows), 2, Inches(6.2), Inches(1.5), Inches(6.6), Inches(4.0)).table
    table.columns[0].width, table.columns[1].width = Inches(3.3), Inches(3.3)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = '\n'.join(split_name(value)) if r and c == 0 else value
            cell.fill.solid()
            cell.fill.fore_color.rgb = TEAL if r == 0 else (LIGHT if r % 2 else WHITE)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(16 if r else 18)
                    run.font.bold = r == 0
                    run.font.color.rgb = WHITE if r == 0 else INK
                    run.font.name = MONO if r and c == 0 else FONT
    text(slide, Inches(6.2), Inches(5.7), Inches(6.6), Inches(0.8),
         'Todas testadas com o resultado exato esperado.', size=20, color=MUTED)
    notes(slide, MA, 35, [
        'Usamos SQLite com migrations, e o create_database.sql cria o banco do zero.',
        'O CRUD é feito pelo ORM nos models. Como o SQLite não tem stored procedures, as consultas especiais ficam nos repositories.',
        'Por exemplo, o ranking de falhas usa JOIN e GROUP BY, e o ranking de custo usa uma CTE. Todas têm teste com o resultado exato.',
        'Agora o João Lucas mostra o README.'])


def technologies(prs):
    slide = base(prs, 'Tecnologias utilizadas', 'Stack', 10)
    cards = [('Frontend', 'JavaScript, Vite 8 e Chart.js'), ('Backend', 'Python 3.12, Flask 3.1 e SQLAlchemy 2.0'),
             ('Banco de dados', 'SQLite e Alembic (migrations)'), ('IA', 'Google Gemini (API REST)'),
             ('Hospedagem', 'PythonAnywhere, com HTTPS'), ('Coletor', 'Python + icmplib; .exe para Windows')]
    for i, (title, body) in enumerate(cards):
        col, row = i % 3, i // 3
        box(slide, Inches(0.6) + col * Inches(4.1), Inches(1.5) + row * Inches(2.6), Inches(3.85), Inches(2.3),
            title, body, color=LIGHT if (col + row) % 2 else MINT, body_size=22, title_size=26)
    notes(slide, FE, 30, [
        'No frontend usamos JavaScript puro com Vite e o Chart.js pros gráficos.',
        'No backend, Python com Flask e SQLAlchemy, e o banco é SQLite com migrations.',
        'A IA é o Gemini do Google, o site está no PythonAnywhere e o coletor é em Python, com um executável pra Windows.',
        'Agora o Matheus mostra como a API é organizada.'])


def readme(prs):
    slide = base(prs, 'README e funcionalidades implementadas', 'Documentação', 14)
    box(slide, Inches(0.6), Inches(1.5), Inches(4.0), Inches(2.4), None, ['27', 'funcionalidades implementadas'],
        color=TEAL, body_color=WHITE, body_size=24, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    slide.shapes[-1].text_frame.paragraphs[0].runs[0].font.size = Pt(66)
    slide.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
    box(slide, Inches(0.6), Inches(4.15), Inches(4.0), Inches(2.4), None, ['6', 'principais, além de CRUD (★)'],
        color=NAVY, body_color=WHITE, body_size=24, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    slide.shapes[-1].text_frame.paragraphs[0].runs[0].font.size = Pt(66)
    slide.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
    text(slide, Inches(5.0), Inches(1.45), Inches(7.8), Inches(0.5), 'As 6 funcionalidades principais (★):', size=22, color=TEAL_DARK, bold=True)
    bullets(slide, Inches(5.0), Inches(2.0), Inches(7.8), Inches(3.4), [
        'Detecção de falhas pelo coletor', 'Diagnóstico com recomendações', 'Prejuízo estimado de cada falha',
        'Explicação com IA', 'Dashboard com ranking de custo', 'Relatório ZIP'], size=22, gap=4)
    text(slide, Inches(5.0), Inches(5.55), Inches(7.8), Inches(1.2),
         'Cada linha do README diz a tela, a rota, o Service e o Repository. Um teste confere se tudo existe no código.',
         size=20, color=MUTED)
    notes(slide, JL, 35, [
        'O README tem a seção Funcionalidades Implementadas, com 27 funcionalidades de ponta a ponta.',
        'Seis delas são as principais, que vão além de CRUD: detecção, diagnóstico, prejuízo, IA, dashboard com ranking e relatório.',
        'Pra cada uma o README diz a tela, a rota, o service e o repository, e um teste automático confere que tudo isso existe.'])


def bonus_quality(prs):
    slide = base(prs, 'Bônus e qualidade', 'Diferenciais', 15)
    cards = [('IA como serviço', ['Gemini explica a falha.', 'Só dados técnicos: nunca nome, e-mail ou IP.']),
             ('Hospedagem', ['PythonAnywhere,', 'com HTTPS.', 'marcelodomingos.', 'pythonanywhere.com']),
             ('Termos de consentimento', ['Aceite obrigatório e com versão registrada.', 'Aviso de Privacidade (LGPD).'])]
    for i, (title, body) in enumerate(cards):
        box(slide, Inches(0.6) + i * Inches(4.1), Inches(1.45), Inches(3.85), Inches(2.85), title, body,
            color=MINT, body_size=20, title_size=24)
    box(slide, Inches(0.6), Inches(4.55), Inches(12.05), Inches(2.25), 'Qualidade e segurança', [
        '186 testes no backend e 20 no frontend, todos passando.',
        'Senhas com hash, sessão protegida contra CSRF, cada empresa só vê os próprios dados.'],
        color=LIGHT, body_size=22, title_size=24)
    notes(slide, JL, 50, [
        'Fizemos os três bônus. A IA como serviço: o Gemini explica a ocorrência, e a gente manda só dados técnicos, nunca nome, e-mail, CNPJ ou IP.',
        'O sistema está hospedado no PythonAnywhere, com HTTPS.',
        'E tem os termos: ninguém usa o sistema sem aceitar, e a versão aceita fica registrada, junto com o aviso de privacidade da LGPD.',
        'Sobre qualidade: são 186 testes automatizados no backend e 20 no frontend, todos passando.',
        'As senhas são guardadas com hash, a sessão é protegida contra CSRF e uma empresa nunca vê os dados de outra.'])


def demo(prs, img):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY
    slide.shapes.add_picture(str(img['logo']), Inches(0.8), Inches(0.8), Inches(1.1), Inches(1.1))
    text(slide, Inches(0.8), Inches(2.2), Inches(11.5), Inches(1.2), 'Demonstração ao vivo', size=54, color=WHITE, bold=True)
    text(slide, Inches(0.8), Inches(3.5), Inches(11.5), Inches(1.6), [
        ('Site: ' + SITE.replace('https://', ''), 26, MINT, True),
        ('Repositório: ' + REPO, 22, WHITE, False)], spacing=1.3)
    text(slide, Inches(0.8), Inches(5.6), Inches(11.5), Inches(1.0), 'Obrigado! Perguntas?', size=34, color=WHITE, bold=True)
    notes(slide, JL, 20, [
        'Agora vamos mostrar o sistema funcionando, ao vivo, no site hospedado (roteiro em apresentacao/ROTEIRO_APRESENTACAO.md).',
        'Depois da demonstração, voltem a este slide para o encerramento: obrigado, e ficamos à disposição para perguntas.'])


def build():
    with tempfile.TemporaryDirectory() as tmp:
        img = prepare_images(Path(tmp))
        prs = Presentation()
        prs.slide_width, prs.slide_height = W, H
        cover(prs, img)
        problem(prs)
        solution(prs)
        use_case(prs, 1, 4, 'Cadastrar dispositivo', 'Entrada de dados',
                 'O técnico cadastra o equipamento que será monitorado: nome, IP, tipo, local e coletor.',
                 img['dispositivo'], 'fluxo-1-cadastrar-dispositivo.png',
                 [('Tela', 'Dispositivos.js'), ('Controller', 'DispositivoController.cadastrar'),
                  ('Service', 'CadastrarDispositivoService'), ('Model', 'Dispositivo.salvar')], MR, 35, [
                     'O primeiro caso de uso é cadastrar um dispositivo, que é entrada de dados.',
                     'O técnico informa nome, IP, tipo, localização e qual coletor vai medir.',
                     'Seguindo o fluxograma: a tela chama o DispositivoController, que chama o CadastrarDispositivoService, que valida e salva pelo model Dispositivo.'])
        use_case(prs, 2, 5, 'Coletor envia medições e o sistema detecta a falha', 'Entrada de dados',
                 'O coletor mede a rede; o servidor abre a ocorrência e faz o diagnóstico sozinho.',
                 img['coletores'], 'fluxo-2-receber-amostras.png',
                 [('Origem', 'Coletor (HTTPS + credencial)'), ('Controller', 'ColetorApiController.amostras'),
                  ('Service', 'RegistrarMedicaoService'), ('Repository', 'FalhaRepository.aberta_do_dispositivo')], MR, 45, [
                     'O segundo caso é o coração do sistema: o coletor envia as medições.',
                     'Ele roda dentro da rede da empresa, mede os equipamentos e manda lotes de amostras com uma credencial própria, que aparece nessa tela de coletores.',
                     'O servidor decide o status: uma medição ruim abre a ocorrência, três seguidas confirmam a queda. Aí ele calcula a severidade e o diagnóstico.',
                     'Agora o Erick mostra a tela da ocorrência.'])
        use_case(prs, 3, 6, 'Ocorrência: diagnóstico, impacto e prejuízo', 'Entrada de dados',
                 'A equipe informa as pessoas afetadas; a severidade e o prejuízo são recalculados na hora.',
                 img['falha'], 'fluxo-3-registrar-impacto.png',
                 [('Tela', 'Falha.js'), ('Controller', 'FalhaController.registrar_impacto'),
                  ('Service', 'RegistrarImpactoService'), ('Repository', 'FalhaRepository.grupo_compartilhado')], ER, 50, [
                     'Esse é o detalhe de uma ocorrência. Aqui aparecem a severidade, o diagnóstico com as causas prováveis e as recomendações.',
                     'O prejuízo estimado mostra a conta aberta: horas de expediente vezes pessoas vezes custo por hora.',
                     'Quando a equipe informa quantas pessoas foram afetadas, o RegistrarImpactoService recalcula a severidade e o prejuízo.',
                     'Se várias falhas acontecem juntas, o repository encontra o grupo pra não contar as mesmas pessoas duas vezes.'])
        use_case(prs, 4, 7, 'Histórico de falhas', 'Recuperação de dados',
                 'Todas as ocorrências da empresa, com filtros por dispositivo, estado, severidade e período.',
                 img['historico'], 'fluxo-4-historico-falhas.png',
                 [('Tela', 'Historico.js'), ('Controller', 'FalhaController.listar'),
                  ('Service', 'ListarFalhasService'), ('Repository', 'FalhaRepository.historico')], ER, 30, [
                     'O quarto caso é de recuperação de dados: o histórico de falhas.',
                     'Dá pra filtrar por dispositivo, estado, severidade e período, com paginação.',
                     'A consulta fica no FalhaRepository.historico, que só traz falhas da empresa logada.'])
        use_case(prs, 5, 8, 'Visão da rede (dashboard)', 'Recuperação de dados',
                 'Estado atual da rede, gráficos de latência e perda e o ranking dos aparelhos que mais custaram.',
                 img['dashboard'], 'fluxo-5-dashboard.png',
                 [('Tela', 'Dashboard.js'), ('Controller', 'DashboardController.obter'),
                  ('Service', 'GerarDashboardService'), ('Repository', 'RankingCustoRepository.top_dispositivos')], FE, 40, [
                     'O quinto caso é a visão da rede, o nosso dashboard.',
                     'Mostra quantos equipamentos estão online, instáveis e offline, os gráficos de latência e perda e o ranking de custo das falhas.',
                     'Os números vêm de consultas SQL nos repositories, como a contagem por status e o top 5 por prejuízo.'])
        architecture(prs)
        technologies(prs)
        directories(prs)
        class_diagram(prs)
        database(prs)
        readme(prs)
        bonus_quality(prs)
        demo(prs, img)
        assert len(prs.slides) == 16
        prs.save(OUT)
    print(f'{OUT.relative_to(ROOT)}: {len(prs.slides)} slides')


if __name__ == '__main__':
    build()
